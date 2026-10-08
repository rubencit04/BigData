import pandas as pd
import sqlite3
import openpyxl
from pathlib import Path

base = Path(r"c:\Users\tteru\OneDrive\Escritorio\Nebrija\Francisco Big Data Aplicado\Data Quality")

controls = [
    ("Facturas", base / "ControlesDQ Nexora (Facturas Completo).xlsx"),
    ("Cobros", base / "ControlesDQ Nexora (Cobros Completo).xlsx"),
    ("Clientes", base / "ControlesDQ Nexora (Clientes Completo).xlsx"),
    ("Empleados", base / "ControlesDQ Nexora (Empleados Completo).xlsx"),
    ("TipoServicio", base / "ControlesDQ Nexora (TipoServicio Completo).xlsx"),
]

ruta_dataset = base / "Dataset_Nexora_DQ.xlsx"
ruta_salida = base / "Resultados_ControlesDQ_Nexora.xlsx"

conn = sqlite3.connect(":memory:")
for sheet in ["Facturas", "Cobros", "Clientes", "Empleados", "TipoServicio"]:
    df = pd.read_excel(ruta_dataset, sheet_name=sheet)
    df.to_sql(sheet, conn, index=False, if_exists="replace")

resultados = []

for entidad, ruta_ctrl in controls:
    wb = openpyxl.load_workbook(ruta_ctrl, data_only=True)
    ws = wb["DQ"]

    for r in range(2, 500):
        control = ws.cell(r, 1).value
        if not control:
            continue

        nivel = ws.cell(r, 2).value
        tabla = ws.cell(r, 3).value
        campo = ws.cell(r, 4).value
        tipo = ws.cell(r, 5).value
        descripcion = ws.cell(r, 6).value
        sql = ws.cell(r, 7).value

        errores = None
        error_sql = ""

        try:
            errores = pd.read_sql_query(sql, conn).iloc[0, 0]
            if pd.isna(errores):
                errores = None
        except Exception as e:
            error_sql = str(e)

        estado = "OK" if errores == 0 else "KO"
        if errores is None:
            estado = "KO"

        resultados.append({
            "EntidadArchivo": entidad,
            "Control": control,
            "Nivel Control": nivel,
            "Tabla": tabla,
            "Campo": campo,
            "Tipo": tipo,
            "Descripcion": descripcion,
            "Errores": errores,
            "Estado": estado,
            "SQL": sql,
            "ErrorSQL": error_sql,
        })

df_resultados = pd.DataFrame(resultados)

resumen = (
    df_resultados.groupby("EntidadArchivo", dropna=False)
    .agg(
        Controles=("Control", "count"),
        OK=("Estado", lambda s: (s == "OK").sum()),
        KO=("Estado", lambda s: (s == "KO").sum()),
        ErroresTotales=("Errores", lambda s: pd.to_numeric(s, errors="coerce").fillna(0).sum()),
    )
    .reset_index()
)

with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
    resumen.to_excel(writer, sheet_name="Resumen", index=False)
    df_resultados.to_excel(writer, sheet_name="Todos", index=False)
    for entidad, _ in controls:
        df_resultados[df_resultados["EntidadArchivo"] == entidad].to_excel(writer, sheet_name=entidad[:31], index=False)

print(f"Resultados exportados a: {ruta_salida}")
